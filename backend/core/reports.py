"""
core/reports.py
---------------
Centralized service for generating PDF and Excel reports for projects, plots, work items, and job items.
"""

import datetime
import io
import os
from decimal import Decimal

from django.http import FileResponse
from rest_framework.renderers import BaseRenderer
from reportlab.lib import colors

class XLSXRenderer(BaseRenderer):
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    format = "xlsx"

    def render(self, data, accepted_media_type=None, renderer_context=None):
        return data


class PDFRenderer(BaseRenderer):
    media_type = "application/pdf"
    format = "pdf"

    def render(self, data, accepted_media_type=None, renderer_context=None):
        return data

from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    Image as PDFImage,
)
import openpyxl


def parse_report_date_range(request):
    """
    Parses start_date and end_date from request parameters.
    Defaults to the current week (Monday through Sunday) if omitted or invalid.
    """
    def parse_date(value):
        if not value:
            return None
        try:
            return datetime.datetime.strptime(str(value).strip(), "%Y-%m-%d").date()
        except (ValueError, TypeError):
            return None

    query_params = getattr(request, "query_params", getattr(request, "GET", {})) if hasattr(request, "query_params") or hasattr(request, "GET") else {}
    start_date = parse_date(query_params.get("start_date"))
    end_date = parse_date(query_params.get("end_date"))

    if not start_date or not end_date:
        today = datetime.date.today()
        weekday = today.weekday()
        start_date = today - datetime.timedelta(days=weekday)
        end_date = start_date + datetime.timedelta(days=6)

    return start_date, end_date


class PDFReportBuilder:
    """
    Builder helper for creating ReportLab PDF documents with consistent styling.
    """

    def __init__(self, page_size=letter, margin=40):
        self.buffer = io.BytesIO()
        self.doc = SimpleDocTemplate(
            self.buffer,
            pagesize=page_size,
            rightMargin=margin,
            leftMargin=margin,
            topMargin=margin,
            bottomMargin=margin,
        )
        self.styles = getSampleStyleSheet()
        self.story = []

    def add_title(self, title_text):
        self.story.append(Paragraph(title_text, self.styles["Title"]))
        self.story.append(Spacer(1, 10))

    def add_heading(self, heading_text, level=2, spacer_after=8):
        style_key = "Heading2" if level == 2 else "Heading1"
        self.story.append(Paragraph(heading_text, self.styles[style_key]))
        if spacer_after > 0:
            self.story.append(Spacer(1, spacer_after))

    def add_metadata_lines(self, lines):
        for line in lines:
            self.story.append(Paragraph(line, self.styles["Normal"]))
        self.story.append(Spacer(1, 15))

    def add_paragraph(self, text, style_name="Normal", spacer_after=8):
        self.story.append(Paragraph(text, self.styles[style_name]))
        if spacer_after > 0:
            self.story.append(Spacer(1, spacer_after))

    def add_spacer(self, height=10):
        self.story.append(Spacer(1, height))

    def add_table(self, data, col_widths=None, header_bg="#f3f4f6", is_summary=False):
        if not data:
            return

        header_style = ParagraphStyle(
            "TableHeader",
            parent=self.styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=9,
            leading=11,
            textColor=colors.black,
            alignment=1 if is_summary else 0,
        )
        header_right_style = ParagraphStyle(
            "TableHeaderRight",
            parent=header_style,
            alignment=2,
        )
        body_style = ParagraphStyle(
            "TableBody",
            parent=self.styles["Normal"],
            fontName="Helvetica",
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#1f2937"),
            alignment=1 if is_summary else 0,
        )
        body_right_style = ParagraphStyle(
            "TableBodyRight",
            parent=body_style,
            alignment=2,
        )

        formatted_data = []
        for row_idx, row in enumerate(data):
            formatted_row = []
            is_header_row = (row_idx == 0)
            num_cols = len(row)

            for col_idx, cell in enumerate(row):
                if isinstance(cell, Paragraph) or hasattr(cell, "draw"):
                    formatted_row.append(cell)
                    continue

                cell_str = str(cell) if cell is not None else "—"
                is_right_col = (not is_summary and col_idx == num_cols - 1)

                if is_header_row:
                    style = header_right_style if is_right_col else header_style
                else:
                    style = body_right_style if is_right_col else body_style

                formatted_row.append(Paragraph(cell_str, style))

            formatted_data.append(formatted_row)

        table = Table(formatted_data, colWidths=col_widths)
        t_style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(header_bg)),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
        ]
        table.setStyle(TableStyle(t_style))
        self.story.append(table)
        self.story.append(Spacer(1, 12))

    def add_photo_figures(self, figures):
        """
        Builds and appends a 2-column table of thumbnail images with captions.
        figures: list of (fig_num, pic_obj, caption_text)
        """
        if not figures:
            return

        caption_style = ParagraphStyle(
            "FigureCaption",
            parent=self.styles.get("Normal", self.styles["BodyText"]),
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#374151"),
            alignment=1,
        )

        fig_table_data = []
        row = []
        for fig_num, pic, caption in figures:
            image_path = getattr(pic.img, 'path', None) if getattr(pic, 'img', None) else None
            if image_path and os.path.exists(image_path):
                try:
                    cell_flowables = [
                        PDFImage(image_path, width=2.4 * inch, height=1.6 * inch),
                        Spacer(1, 4),
                        Paragraph(f"<b>Fig. {fig_num}</b>: {caption}", caption_style)
                    ]
                    row.append(cell_flowables)
                    if len(row) == 2:
                        fig_table_data.append(row)
                        row = []
                except Exception:
                    continue

        if row:
            while len(row) < 2:
                row.append("")
            fig_table_data.append(row)

        if not fig_table_data:
            return

        figures_table = Table(fig_table_data, colWidths=[265, 265])
        figures_table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ]))
        self.story.append(figures_table)

    def build_response(self, filename):
        self.doc.build(self.story)
        self.buffer.seek(0)
        return FileResponse(
            self.buffer,
            as_attachment=True,
            filename=filename,
            content_type="application/pdf",
        )


class ExcelReportBuilder:
    """
    Builder helper for creating OpenPyXL spreadsheets with consistent styling.
    """

    def __init__(self, title="Summary"):
        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active
        self.ws.title = title

    def add_header(self, title, metadata_pairs):
        self.ws.append([title])
        for label, val in metadata_pairs:
            if isinstance(val, (list, tuple)):
                self.ws.append([label, *val])
            else:
                self.ws.append([label, val])
        self.ws.append([])

    def add_table(self, sheet_title, headers, rows):
        target_ws = self.wb.create_sheet(title=sheet_title) if sheet_title else self.ws
        if headers:
            target_ws.append(headers)
        for row in rows:
            target_ws.append(row)

    def build_response(self, filename):
        buffer = io.BytesIO()
        self.wb.save(buffer)
        buffer.seek(0)
        return FileResponse(
            buffer,
            as_attachment=True,
            filename=filename,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )


def build_progress_report_pdf(title, metadata_lines, reports_qs, filename, entity_level="project"):
    """
    Generic progress report builder for project, plot, work item, or job item reports.
    Renders job reports in a structured table format with columns: Date, Job, % Compl, Notes, Issues.
    """
    builder = PDFReportBuilder()
    builder.add_title(title)
    builder.add_metadata_lines(metadata_lines)

    if not reports_qs.exists():
        builder.add_paragraph("No job reports submitted for the selected criteria in this period.")
        return builder.build_response(filename)

    figures = []
    fig_counter = 1

    headers = ["Date", "Job", "% Compl", "Notes", "Issues"]
    rows = []

    for rep in reports_qs:
        r_date = rep.report_date.isoformat() if hasattr(rep.report_date, "isoformat") else str(rep.report_date)
        notes_text = rep.notes or "—"
        issues_text = rep.issues_encountered or "—"
        rows.append([
            r_date,
            rep.job_item.job_name,
            f"{rep.percentage_job_progress}%",
            notes_text,
            issues_text,
        ])

        for pic in rep.photos.all():
            cap = f"{rep.job_item.job_name} ({r_date})"
            pic_cap = getattr(pic, "description", None) or getattr(pic, "caption", None)
            if pic_cap:
                cap += f" — {pic_cap}"
            figures.append((fig_counter, pic, cap))
            fig_counter += 1

    builder.add_table([headers, *rows], col_widths=[65, 115, 55, 160, 137])

    if figures:
        builder.add_spacer(10)
        builder.add_heading("Attached Photographic Figures", level=2)
        builder.add_paragraph("Catalog of job progress photos referenced in the report above:")
        builder.add_spacer(8)
        builder.add_photo_figures(figures)

    return builder.build_response(filename)


def build_financial_report_pdf(
    title,
    metadata_lines,
    summary_headers,
    summary_values,
    summary_col_widths,
    breakdown_heading=None,
    breakdown_headers=None,
    breakdown_rows=None,
    breakdown_col_widths=None,
    itemized_heading=None,
    itemized_headers=None,
    itemized_rows=None,
    itemized_col_widths=None,
    filename="financial_report.pdf",
    breakdowns=None,
):
    """
    Generic financial report builder for PDF format.
    Supports multiple breakdown sections via `breakdowns`:
    breakdowns = [(heading, headers, rows, col_widths), ...]
    """
    builder = PDFReportBuilder()
    builder.add_title(title)
    builder.add_metadata_lines(metadata_lines)

    # Summary table
    if summary_headers and summary_values:
        builder.add_table([summary_headers, summary_values], col_widths=summary_col_widths, is_summary=True)

    # Multiple Breakdown tables
    if breakdowns:
        for heading, headers, rows, col_widths in breakdowns:
            if heading and headers and rows:
                builder.add_heading(heading, level=2)
                builder.add_table([headers, *rows], col_widths=col_widths)
    elif breakdown_heading and breakdown_headers and breakdown_rows:
        builder.add_heading(breakdown_heading, level=2)
        builder.add_table([breakdown_headers, *breakdown_rows], col_widths=breakdown_col_widths)

    # Itemized expenses table
    if itemized_heading:
        builder.add_heading(itemized_heading, level=2)
        if not itemized_rows:
            builder.add_paragraph("No expenses recorded for this entity.")
        else:
            builder.add_table([itemized_headers, *itemized_rows], col_widths=itemized_col_widths)

    return builder.build_response(filename)


def build_financial_report_excel(
    header_title,
    metadata_pairs,
    summary_headers,
    summary_values,
    breakdown_title=None,
    breakdown_headers=None,
    breakdown_rows=None,
    itemized_sheet_title=None,
    itemized_headers=None,
    itemized_rows=None,
    filename="financial_report.xlsx",
    breakdowns=None,
):
    """
    Generic financial report builder for Excel XLSX format.
    Supports multiple breakdown sections via `breakdowns`:
    breakdowns = [(title, headers, rows), ...]
    """
    builder = ExcelReportBuilder(title="Financial Summary")
    builder.add_header(header_title, metadata_pairs)

    # Summary
    if summary_headers and summary_values:
        builder.ws.append(summary_headers)
        builder.ws.append(summary_values)
        builder.ws.append([])

    # Multiple Breakdown tables
    if breakdowns:
        for b_title, b_headers, b_rows in breakdowns:
            if b_title and b_rows:
                builder.ws.append([b_title])
                if b_headers:
                    builder.ws.append(b_headers)
                for r in b_rows:
                    builder.ws.append(r)
                builder.ws.append([])
    elif breakdown_title and breakdown_rows:
        builder.ws.append([breakdown_title])
        if breakdown_headers:
            builder.ws.append(breakdown_headers)
        for r in breakdown_rows:
            builder.ws.append(r)
        builder.ws.append([])

    # Itemized Expenses (rendered on primary sheet for instant visibility)
    if itemized_headers:
        sheet_label = itemized_sheet_title or "Itemized Expenditures"
        builder.ws.append([sheet_label])
        builder.ws.append(itemized_headers)
        if not itemized_rows:
            builder.ws.append(["No expenses recorded for this entity."])
        else:
            for r in itemized_rows:
                builder.ws.append(r)
        builder.ws.append([])

    # Also add a dedicated tab for itemized expenses if requested
    if itemized_sheet_title and itemized_headers and itemized_rows:
        builder.add_table(itemized_sheet_title, itemized_headers, itemized_rows)

    return builder.build_response(filename)


def build_progress_report_excel(title, metadata_lines, reports_qs, filename):
    """
    Generic daily progress report builder for Excel XLSX format.
    """
    builder = ExcelReportBuilder(title="Progress Summary")
    metadata_pairs = []
    for line in metadata_lines:
        if "|" in line:
            parts = line.split("|")
            for p in parts:
                if ":" in p:
                    k, v = p.split(":", 1)
                    metadata_pairs.append((k.strip() + ":", v.strip()))
                else:
                    metadata_pairs.append(("Info:", p.strip()))
        elif ":" in line:
            k, v = line.split(":", 1)
            metadata_pairs.append((k.strip() + ":", v.strip()))
        else:
            metadata_pairs.append(("Info:", line))

    headers = ["Date", "Job", "% Compl", "Notes", "Issues"]
    rows = []
    for rep in reports_qs:
        r_date = rep.report_date.isoformat() if hasattr(rep.report_date, "isoformat") else str(rep.report_date)
        rows.append([
            r_date,
            rep.job_item.job_name,
            f"{rep.percentage_job_progress}%",
            rep.notes or "—",
            rep.issues_encountered or "—"
        ])

    builder.add_header(title, metadata_pairs)
    builder.add_table("Job Reports", headers, rows)
    return builder.build_response(filename)

