import pytest
import os
from PyPDF2 import PdfReader
from openpyxl import load_workbook
from playwright.sync_api import Page, expect
from utils.factories import ProjectFactory, PlotFactory, WorkFactory, JobFactory, UserFactory
from backend.core.models import Expense

pytestmark = pytest.mark.reports_export

@pytest.fixture
def complex_project(project_creator):
    project = ProjectFactory.create(name="Finance Export Project", creator=project_creator, budget=500000)
    plot1 = PlotFactory.create(project=project, name="Plot 1", budget=200000)
    plot2 = PlotFactory.create(project=project, name="Plot 2", budget=300000)
    work1 = WorkFactory.create(plot=plot1, name="Work 1", budget=100000)
    job1 = JobFactory.create(work=work1, name="Job 1", budget=50000)
    job2 = JobFactory.create(work=work1, name="Job 2", budget=50000)
    
    # Expenses
    Expense.objects.create(job_item=job1, amount=10000, description="Materials 1")
    Expense.objects.create(job_item=job1, amount=20000, description="Materials 2")
    Expense.objects.create(job_item=job2, amount=40000, description="Equipment")
    return project

def test_financial_report_generation_project_level(auth_page_creator: Page, base_url: str, complex_project):
    page = auth_page_creator
    page.goto(f"{base_url}/projects/{complex_project.id}")
    
    with page.expect_download() as download_info:
        page.get_by_role("button", name="Export").click()
        page.get_by_role("menuitem", name="Financial Report (PDF)").click()
    
    download = download_info.value
    path = download.path()
    assert download.suggested_filename.endswith(".pdf")
    
    reader = PdfReader(path)
    text = ""
    for p in reader.pages:
        text += p.extract_text()
    
    # Assert data hierarchy and aggregation is in the text
    assert "Finance Export Project" in text
    assert "Plot 1" in text
    assert "Materials 1" in text
    assert "10,000" in text

def test_financial_report_export_xlsx(auth_page_creator: Page, base_url: str, complex_project):
    page = auth_page_creator
    page.goto(f"{base_url}/projects/{complex_project.id}")
    
    with page.expect_download() as download_info:
        page.get_by_role("button", name="Export").click()
        page.get_by_role("menuitem", name="Financial Report (Excel)").click()
    
    download = download_info.value
    path = download.path()
    assert download.suggested_filename.endswith(".xlsx")
    
    wb = load_workbook(path)
    sheet = wb.active
    
    # Simple check for values
    cell_values = [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value]
    joined_values = " ".join(cell_values)
    
    assert "Finance Export Project" in joined_values
    assert "Materials 1" in joined_values

def test_financial_report_plot_level(auth_page_creator: Page, base_url: str, complex_project):
    page = auth_page_creator
    plot = complex_project.plots.first()
    page.goto(f"{base_url}/plots/{plot.id}")
    
    with page.expect_download() as download_info:
        page.get_by_role("button", name="Export").click()
        page.get_by_role("menuitem", name="Financial Report (PDF)").click()
        
    download = download_info.value
    path = download.path()
    reader = PdfReader(path)
    text = ""
    for p in reader.pages:
        text += p.extract_text()
        
    assert "Plot 1" in text
    assert "Plot 2" not in text

def test_financial_report_work_level(auth_page_creator: Page, base_url: str, complex_project):
    page = auth_page_creator
    work = complex_project.plots.first().works.first()
    page.goto(f"{base_url}/works/{work.id}")
    
    with page.expect_download() as download_info:
        page.get_by_role("button", name="Export").click()
        page.get_by_role("menuitem", name="Financial Report (PDF)").click()
        
    download = download_info.value
    path = download.path()
    reader = PdfReader(path)
    text = ""
    for p in reader.pages:
        text += p.extract_text()
        
    assert "Work 1" in text
    assert "Job 1" in text
    assert "Job 2" in text

def test_financial_report_job_level(auth_page_creator: Page, base_url: str, complex_project):
    page = auth_page_creator
    job = complex_project.plots.first().works.first().jobs.first()
    page.goto(f"{base_url}/jobs/{job.id}")
    
    with page.expect_download() as download_info:
        page.get_by_role("button", name="Export").click()
        page.get_by_role("menuitem", name="Financial Report (PDF)").click()
        
    download = download_info.value
    path = download.path()
    reader = PdfReader(path)
    text = ""
    for p in reader.pages:
        text += p.extract_text()
        
    assert "Job 1" in text
    assert "Materials 1" in text
    assert "Materials 2" in text

def test_reportlab_escaping_special_characters(auth_page_creator: Page, base_url: str, project_creator):
    name = "Project & < > \" ' test"
    project = ProjectFactory.create(name=name, creator=project_creator)
    plot = PlotFactory.create(project=project, name="Plot & Co.")
    work = WorkFactory.create(plot=plot, name="Work <tag>")
    job = JobFactory.create(work=work, name="Job 'quote'")
    Expense.objects.create(job_item=job, amount=100, description="&<>\"'")
    
    page = auth_page_creator
    page.goto(f"{base_url}/projects/{project.id}")
    
    with page.expect_download() as download_info:
        page.get_by_role("button", name="Export").click()
        page.get_by_role("menuitem", name="Financial Report (PDF)").click()
        
    download = download_info.value
    path = download.path()
    
    reader = PdfReader(path)
    text = ""
    for p in reader.pages:
        text += p.extract_text()
        
    assert "Project & < > \" ' test" in text
    assert "&<>\"'" in text
